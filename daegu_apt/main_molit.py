"""
국토교통부 실거래가 독립 수집기 — Naver 의존 없음

사용법:
  python main_molit.py --molit-key YOUR_KEY
  python main_molit.py --molit-key YOUR_KEY --district 수성구
  python main_molit.py --molit-key YOUR_KEY --months 6
  python main_molit.py --molit-key YOUR_KEY --district 달서구 --output 달서구_실거래.xlsx

출력 컬럼:
  구 | 법정동 | 도로명 | 아파트명 | 건축년도 | 전용면적㎡ | 층
  | 거래금액(만원) | 거래유형 | 거래년도 | 거래월 | 거래일 | 등기일자
  | 해제여부 | 해제사유발생일
  | 근처지하철 | 지하철직선도보(분) | 위도 | 경도
"""
import argparse
import os
import sys
from datetime import datetime
from pathlib import Path


def _load_env():
    env_path = Path(__file__).parent / '.env'
    if not env_path.exists():
        return
    with open(env_path, encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith('#') or '=' not in line:
                continue
            key, _, val = line.partition('=')
            key = key.strip()
            val = val.strip().strip('"').strip("'")
            if key and key not in os.environ:
                os.environ[key] = val


_load_env()

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


HEADER_PURPLE = PatternFill('solid', fgColor='403151')
WHITE_FONT    = Font(color='FFFFFF', bold=True, size=10)

NUM_COLS = ['거래금액', '전용면적', '지하철직선도보', '위도', '경도']


def _style_sheet(ws):
    for cell in ws[1]:
        cell.fill      = HEADER_PURPLE
        cell.font      = WHITE_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len    = max((len(str(c.value or '')) for c in col), default=8)
        ws.column_dimensions[col_letter].width = min(max_len + 4, 30)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            col_name = ws.cell(1, cell.column).value or ''
            if any(k in col_name for k in NUM_COLS):
                cell.number_format = '#,##0'
                cell.alignment     = Alignment(horizontal='right')

    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 22


def main():
    parser = argparse.ArgumentParser(
        description='국토교통부 실거래가 수집기 (Naver 불필요)',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    parser.add_argument('--molit-key', default=os.environ.get('MOLIT_API_KEY', ''),
                        help='국토교통부 API 키 (data.go.kr 발급)')
    parser.add_argument('--district', default='all',
                        choices=['all', '중구', '동구', '서구', '남구', '북구',
                                 '수성구', '달서구', '달성군'],
                        help='수집 지역 (기본: all)')
    parser.add_argument('--months', type=int, default=3,
                        help='최근 N개월 (기본: 3)')
    parser.add_argument('--output', default='',
                        help='출력 파일명 (기본: {지역}_실거래가_YYYYMMDD_HHMMSS.xlsx)')
    args = parser.parse_args()

    if not args.molit_key:
        print('[오류] --molit-key 또는 환경변수 MOLIT_API_KEY가 필요합니다.')
        print('       발급: https://www.data.go.kr → "아파트매매 실거래가 상세자료"')
        sys.exit(1)

    ts         = datetime.now().strftime('%Y%m%d_%H%M%S')
    region_tag = args.district if args.district != 'all' else '대구전체'
    out_path   = args.output or f'{region_tag}_실거래가_{ts}.xlsx'

    print(f'=== 국토교통부 실거래가 수집 [{region_tag}] 최근 {args.months}개월 ===')

    from molit_scraper import MolitScraper
    df = MolitScraper(args.molit_key).collect(district=args.district, months=args.months)

    if df.empty:
        print('수집된 데이터가 없습니다.')
        sys.exit(1)

    # 취소 거래 표시 (해제여부 = 'O' 인 행)
    cancelled = (df['해제여부'] == 'O').sum()
    if cancelled:
        print(f'  ※ 취소된 거래 {cancelled}건 포함 (해제여부=O)')

    df_sorted = df.sort_values(
        ['거래년도', '거래월', '거래일'], ascending=False
    )

    with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
        df_sorted.to_excel(writer, sheet_name='실거래가', index=False)
        _style_sheet(writer.sheets['실거래가'])

    print(f'\n--- 수집 요약 ---')
    print(f'  총 {len(df):,}건')
    print(f'  기간: {df["거래년도"].min()}.{df["거래월"].min()} ~ {df["거래년도"].max()}.{df["거래월"].max()}')
    print(f'  파일: {os.path.abspath(out_path)}')


if __name__ == '__main__':
    main()
